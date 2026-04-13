#!/usr/bin/env python3
"""
WAVES PPT Generator
Academic style: navy header, clean typography, concise bullet points
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import re
import os

# Academic color scheme
HEADER_COLOR = RGBColor(0, 51, 102)  # Navy blue
ACCENT_COLOR = RGBColor(70, 130, 180)  # Steel blue
TEXT_COLOR = RGBColor(33, 33, 33)  # Dark gray
BG_COLOR = RGBColor(255, 255, 255)  # White

def parse_slides(filepath):
    """Parse slides from merged text file."""
    with open(filepath, 'r', encoding='utf-8') as f:
        content = f.read()
    
    # Split by SLIDE markers
    slides = []
    parts = re.split(r'\n=+\nSLIDE (\d+[a-z]*)\s*—\s*([^\n]+)\n', content)
    
    # parts[0] is before first slide, then alternating: num, title, content
    i = 1
    while i < len(parts):
        if i + 2 < len(parts):
            slide_num = parts[i]
            slide_title = parts[i+1].strip()
            slide_content = parts[i+2].strip()
            slides.append({
                'num': slide_num,
                'title': slide_title,
                'content': slide_content
            })
        i += 3
    
    return slides

def create_presentation(slides, output_path):
    """Create PowerPoint presentation."""
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Blank layout
    blank_layout = prs.slide_layouts[6]
    
    for idx, slide_data in enumerate(slides):
        slide = prs.slides.add_slide(blank_layout)
        
        # Title slide (first slide)
        if idx == 0:
            create_title_slide(slide, slide_data)
            continue
        
        # Content slides
        create_content_slide(slide, slide_data, idx)
    
    prs.save(output_path)
    print(f"Saved: {output_path}")

def create_title_slide(slide, data):
    """Create title slide."""
    # Navy blue background rectangle at top
    header = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(2.5)
    )
    header.fill.solid()
    header.fill.fore_color.rgb = HEADER_COLOR
    header.line.fill.background()
    
    # Title text
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(12.333), Inches(1))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "WAVES"
    p.font.size = Pt(54)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.alignment = PP_ALIGN.CENTER
    
    # Subtitle
    p2 = tf.add_paragraph()
    p2.text = "Window-based Adaptive Violation Detection with Elastic Streaming"
    p2.font.size = Pt(24)
    p2.font.color.rgb = RGBColor(200, 220, 240)
    p2.alignment = PP_ALIGN.CENTER
    
    # Content area
    content_box = slide.shapes.add_textbox(Inches(0.75), Inches(3), Inches(11.833), Inches(4))
    tf = content_box.text_frame
    tf.word_wrap = True
    
    # Parse content lines
    lines = data['content'].strip().split('\n')
    for i, line in enumerate(lines[:8]):  # Limit to 8 lines
        line = line.strip()
        if not line or line.startswith('==='):
            continue
        
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        
        # Remove markdown markers
        line = line.replace('**', '')
        p.text = line
        p.font.size = Pt(18)
        p.font.color.rgb = TEXT_COLOR
        
        if line.startswith('- '):
            p.level = 0
        elif line.startswith('  - '):
            p.level = 1
            p.font.size = Pt(16)
        
        p.space_after = Pt(6)

def create_content_slide(slide, data, idx):
    """Create content slide."""
    # Header bar
    header = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.2)
    )
    header.fill.solid()
    header.fill.fore_color.rgb = HEADER_COLOR
    header.line.fill.background()
    
    # Slide number and title
    num_box = slide.shapes.add_textbox(Inches(0.4), Inches(0.25), Inches(1), Inches(0.7))
    tf = num_box.text_frame
    p = tf.paragraphs[0]
    p.text = f"Slide {data['num']}"
    p.font.size = Pt(14)
    p.font.color.rgb = RGBColor(180, 200, 220)
    p.font.bold = False
    
    # Title
    title_box = slide.shapes.add_textbox(Inches(1.5), Inches(0.2), Inches(11), Inches(0.8))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = data['title']
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    
    # Content area
    content_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(12.333), Inches(5.7))
    tf = content_box.text_frame
    tf.word_wrap = True
    
    # Parse and simplify content
    content = simplify_content(data['content'])
    lines = content.strip().split('\n')
    
    first_para = True
    for line in lines[:25]:  # Limit to 25 lines for readability
        line = line.strip()
        if not line or line.startswith('===') or line.startswith('---'):
            continue
        
        # Skip code blocks
        if line.startswith('```') or line.startswith('    '):
            continue
        
        if first_para:
            p = tf.paragraphs[0]
            first_para = False
        else:
            p = tf.add_paragraph()
        
        # Clean and format
        line = clean_line(line)
        p.text = line
        p.font.size = Pt(14)
        p.font.color.rgb = TEXT_COLOR
        p.space_after = Pt(4)
        
        # Handle bullet levels
        if line.startswith('- '):
            p.text = line[2:]
            p.level = 0
        elif line.startswith('  - '):
            p.text = line[4:]
            p.level = 1
            p.font.size = Pt(12)
        elif line.startswith('    '):
            p.text = line[4:]
            p.level = 2
            p.font.size = Pt(11)
        elif line.startswith('| '):
            p.text = line[2:]
            p.font.size = Pt(10)
            p.font.name = 'Courier New'
        elif '→' in line:
            p.font.color.rgb = ACCENT_COLOR
        elif line.startswith('**'):
            p.font.bold = True
            p.text = line.replace('**', '')

def clean_line(line):
    """Clean a line for display."""
    # Remove markdown formatting
    line = line.replace('**', '')
    line = line.replace('*', '')
    line = line.replace('`', '')
    
    # Truncate very long lines
    if len(line) > 100:
        line = line[:97] + '...'
    
    return line

def simplify_content(content):
    """Simplify slide content for PPT display."""
    lines = content.strip().split('\n')
    simplified = []
    
    skip_patterns = [
        'Source:', 'Docs:', 'GitHub:', 'Blog:', 'DOI:',
        'arXiv:', 'Nguồn:'
    ]
    
    for line in lines:
        line = line.strip()
        
        # Skip URLs and references
        if any(line.startswith(p) or line.startswith(p.replace(':', ': ')) for p in skip_patterns):
            continue
        if line.startswith('https://') or line.startswith('http://'):
            continue
        
        # Skip very long lines with many details
        if line.startswith('  Source:'):
            continue
        
        simplified.append(line)
    
    return '\n'.join(simplified)

def main():
    input_file = 'slides_merged.txt'
    output_file = 'WAVES_Presentation.pptx'
    
    if not os.path.exists(input_file):
        print(f"Error: {input_file} not found")
        return
    
    print(f"Parsing slides from {input_file}...")
    slides = parse_slides(input_file)
    print(f"Found {len(slides)} slides")
    
    print(f"Creating presentation...")
    create_presentation(slides, output_file)
    print("Done!")

if __name__ == '__main__':
    main()
